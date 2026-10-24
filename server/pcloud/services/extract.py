"""Text extraction from documents (txt/code, PDF, DOCX, XLSX/CSV)."""
from __future__ import annotations

import csv
import io
import os

TEXT_EXTS = {
    ".txt", ".md", ".markdown", ".csv", ".tsv", ".json", ".xml", ".html", ".htm", ".css", ".js", ".ts", ".tsx",
    ".jsx", ".py", ".java", ".c", ".cpp", ".h", ".cs", ".go", ".rs", ".rb", ".php", ".sh", ".ps1", ".bat", ".sql",
    ".yml", ".yaml", ".toml", ".ini", ".cfg", ".conf", ".log", ".rst", ".tex", ".srt", ".vtt",
}
DOC_EXTS = {".pdf", ".docx", ".xlsx", ".xlsm"}
EXTRACTABLE = TEXT_EXTS | DOC_EXTS


def can_extract(name: str) -> bool:
    return os.path.splitext(name)[1].lower() in EXTRACTABLE


def _read_text(path: str, limit: int) -> str:
    with open(path, "rb") as f:
        raw = f.read(limit)
    for enc in ("utf-8-sig", "utf-16") if raw[:2] in (b"\xff\xfe", b"\xfe\xff") else ("utf-8-sig", "cp1252"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


def extract_text(path: str, max_chars: int = 200_000) -> str:
    ext = os.path.splitext(path)[1].lower()
    if ext in TEXT_EXTS:
        return _read_text(path, max_chars * 4)[:max_chars]
    if ext == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(path)
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception:
                return "[encrypted PDF - text not available]"
        parts, total = [], 0
        for i, page in enumerate(reader.pages):
            t = page.extract_text() or ""
            parts.append(f"\n--- page {i + 1} ---\n{t}")
            total += len(t)
            if total >= max_chars:
                break
        return "".join(parts)[:max_chars]
    if ext == ".docx":
        import docx
        d = docx.Document(path)
        parts = [p.text for p in d.paragraphs]
        for table in d.tables:
            for row in table.rows:
                parts.append(" | ".join(c.text for c in row.cells))
        return "\n".join(parts)[:max_chars]
    if ext in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        out = io.StringIO()
        for ws in wb.worksheets:
            out.write(f"\n--- sheet: {ws.title} ---\n")
            w = csv.writer(out)
            for row in ws.iter_rows(values_only=True):
                w.writerow(["" if v is None else v for v in row])
                if out.tell() > max_chars:
                    break
        wb.close()
        return out.getvalue()[:max_chars]
    raise ValueError(f"Text extraction is not supported for {ext or 'this file type'}")


def analyze_table(path: str, sheet: str | None = None, max_rows: int = 100_000) -> dict:
    """Summarize a spreadsheet/CSV: columns, types, numeric stats, sample rows."""
    ext = os.path.splitext(path)[1].lower()
    rows: list[list] = []
    sheets: list[str] = []
    if ext in (".csv", ".tsv"):
        text = _read_text(path, 200 * 1024 * 1024)
        dialect = csv.excel_tab if ext == ".tsv" else csv.excel
        try:
            dialect = csv.Sniffer().sniff(text[:4096])
        except csv.Error:
            pass
        for i, r in enumerate(csv.reader(io.StringIO(text), dialect)):
            rows.append(r)
            if i >= max_rows:
                break
    elif ext in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        sheets = wb.sheetnames
        ws = wb[sheet] if sheet and sheet in sheets else wb.worksheets[0]
        sheet = ws.title
        for i, r in enumerate(ws.iter_rows(values_only=True)):
            rows.append(list(r))
            if i >= max_rows:
                break
        wb.close()
    else:
        raise ValueError("Only CSV, TSV and XLSX files can be analyzed")
    rows = [r for r in rows if any(v not in (None, "") for v in r)]
    if not rows:
        return {"sheet": sheet, "sheets": sheets, "rows": 0, "columns": []}
    header = [str(h) if h not in (None, "") else f"column_{i + 1}" for i, h in enumerate(rows[0])]
    body = rows[1:]
    columns = []
    for ci, name in enumerate(header):
        values = [r[ci] if ci < len(r) else None for r in body]
        nums = []
        for v in values:
            if isinstance(v, (int, float)) and not isinstance(v, bool):
                nums.append(float(v))
            elif isinstance(v, str):
                try:
                    nums.append(float(v.replace(",", "").replace("$", "").replace("%", "").strip()))
                except ValueError:
                    pass
        non_empty = [v for v in values if v not in (None, "")]
        col = {"name": name, "non_empty": len(non_empty), "distinct": len({str(v) for v in non_empty})}
        if non_empty and len(nums) >= 0.8 * len(non_empty):
            col.update(type="number", sum=sum(nums), min=min(nums), max=max(nums), mean=sum(nums) / len(nums))
        else:
            col.update(type="text", examples=[str(v) for v in non_empty[:3]])
        columns.append(col)
    return {"sheet": sheet, "sheets": sheets, "rows": len(body), "columns": columns,
            "sample": [[("" if v is None else str(v)) for v in r] for r in body[:10]]}
