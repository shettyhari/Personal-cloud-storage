"""Built-in, pre-approved background tasks. None of these execute shell commands."""
from __future__ import annotations

import csv
import fnmatch
import hashlib
import os
import tarfile
import time
import zipfile
from collections import defaultdict
from datetime import datetime

from ..paths import PathError, get_root, resolve, validate_name
from . import extract, fileops, index
from .jobs import JobContext, handler


def _out_name(directory: str, name: str) -> str:
    return fileops._unique_name(directory, validate_name(name))


def _walk_files(abs_path: str):
    if os.path.isfile(abs_path):
        yield abs_path
        return
    for dirpath, dirnames, files in os.walk(abs_path):
        dirnames[:] = [d for d in dirnames if d not in index.SKIP_DIRS]
        for f in files:
            yield os.path.join(dirpath, f)


@handler("index_scan", "Rebuild search index", resumable=True)
def job_index(ctx: JobContext, params: dict):
    return index.scan_all(progress=lambda f, m: ctx.progress(f, m), cancelled=ctx.cancelled)


@handler("compress", "Compress to ZIP")
def job_compress(ctx: JobContext, p: dict):
    root = get_root(int(p["root_id"]))
    dest_abs, dest_rel = resolve(root, p.get("dest_dir", ""), must_exist=True)
    name = _out_name(dest_abs, p.get("archive_name") or f"archive-{datetime.now():%Y%m%d-%H%M%S}.zip")
    if not name.lower().endswith(".zip"):
        name += ".zip"
    sources = [resolve(root, rel, must_exist=True) for rel in p["paths"]]
    files = [(f, src_abs) for src_abs, _ in sources for f in _walk_files(src_abs)]
    total = sum(os.path.getsize(f) for f, _ in files) or 1
    done = 0
    out = os.path.join(dest_abs, name)
    tmp = out + ".part"
    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED, allowZip64=True) as z:
        for f, src_abs in files:
            ctx.check()
            arc = os.path.relpath(f, os.path.dirname(src_abs))
            z.write(f, arc)
            done += os.path.getsize(f)
            ctx.progress(done / total, f"Adding {arc}")
    os.replace(tmp, out)
    rel = f"{dest_rel}/{name}" if dest_rel else name
    index.upsert(root, out, rel)
    ctx.log(f"wrote {rel} ({len(files)} files)")
    return {"output": rel, "files": len(files), "size": os.path.getsize(out)}


@handler("extract", "Extract archive")
def job_extract(ctx: JobContext, p: dict):
    root = get_root(int(p["root_id"]))
    src_abs, _ = resolve(root, p["path"], must_exist=True)
    dest_abs, dest_rel = resolve(root, p.get("dest_dir", ""), must_exist=True)
    folder = _out_name(dest_abs, os.path.splitext(os.path.basename(src_abs))[0] or "extracted")
    target = os.path.join(dest_abs, folder)
    os.makedirs(target)
    real_target = os.path.realpath(target)

    def safe_member(name: str) -> str:
        dest = os.path.realpath(os.path.join(target, name))
        if os.path.commonpath([os.path.normcase(dest), os.path.normcase(real_target)]) != os.path.normcase(real_target):
            raise PathError(f"Archive entry escapes the destination: {name}")
        return dest

    count = 0
    if zipfile.is_zipfile(src_abs):
        with zipfile.ZipFile(src_abs) as z:
            members = z.infolist()
            for i, m in enumerate(members):
                ctx.check()
                safe_member(m.filename)
                z.extract(m, target)
                count += 1
                ctx.progress((i + 1) / len(members), f"Extracting {m.filename}")
    elif tarfile.is_tarfile(src_abs):
        with tarfile.open(src_abs) as t:
            members = t.getmembers()
            for i, m in enumerate(members):
                ctx.check()
                safe_member(m.name)
                t.extract(m, target, filter="data")
                count += 1
                ctx.progress((i + 1) / len(members), f"Extracting {m.name}")
    else:
        raise ValueError("Only ZIP and TAR archives are supported")
    rel = f"{dest_rel}/{folder}" if dest_rel else folder
    index.upsert(root, target, rel, recursive=True)
    return {"output": rel, "entries": count}


@handler("inventory", "Generate file inventory")
def job_inventory(ctx: JobContext, p: dict):
    root = get_root(int(p["root_id"]))
    src_abs, src_rel = resolve(root, p.get("path", ""), must_exist=True)
    dest_abs, dest_rel = resolve(root, p.get("dest_dir", src_rel), must_exist=True)
    name = _out_name(dest_abs, f"inventory-{datetime.now():%Y%m%d-%H%M%S}.csv")
    out = os.path.join(dest_abs, name)
    n = 0
    with open(out, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(["path", "name", "extension", "size_bytes", "modified", "created"])
        for fp in _walk_files(src_abs):
            ctx.check()
            if os.path.normcase(fp) == os.path.normcase(out):
                continue
            try:
                st = os.stat(fp)
            except OSError:
                continue
            w.writerow([os.path.relpath(fp, root.real).replace("\\", "/"), os.path.basename(fp),
                        os.path.splitext(fp)[1].lower(), st.st_size, fileops.fmt_time(st.st_mtime),
                        fileops.fmt_time(st.st_ctime)])
            n += 1
            if n % 500 == 0:
                ctx.progress(None, f"{n} files listed")
    rel = f"{dest_rel}/{name}" if dest_rel else name
    index.upsert(root, out, rel)
    return {"output": rel, "files": n}


@handler("search_dir", "Search directory")
def job_search(ctx: JobContext, p: dict):
    root = get_root(int(p["root_id"]))
    src_abs, _ = resolve(root, p.get("path", ""), must_exist=True)
    pattern = p.get("pattern") or "*"
    contains = (p.get("contains") or "").lower()
    hits = []
    scanned = 0
    for fp in _walk_files(src_abs):
        ctx.check()
        scanned += 1
        if scanned % 1000 == 0:
            ctx.progress(None, f"Scanned {scanned} files, {len(hits)} matches")
        if not fnmatch.fnmatch(os.path.basename(fp).lower(), pattern.lower()):
            continue
        if contains:
            if not extract.can_extract(fp) or os.path.getsize(fp) > 20 * 1024 * 1024:
                continue
            try:
                if contains not in extract.extract_text(fp, 2_000_000).lower():
                    continue
            except Exception:
                continue
        hits.append({"path": os.path.relpath(fp, root.real).replace("\\", "/"), "size": os.path.getsize(fp)})
        if len(hits) >= 5000:
            break
    return {"scanned": scanned, "matches": hits}


@handler("checksums", "Generate checksums")
def job_checksums(ctx: JobContext, p: dict):
    root = get_root(int(p["root_id"]))
    src_abs, src_rel = resolve(root, p.get("path", ""), must_exist=True)
    files = list(_walk_files(src_abs))
    results = {}
    for i, fp in enumerate(files):
        ctx.check()
        results[os.path.relpath(fp, root.real).replace("\\", "/")] = fileops.checksum(fp)
        ctx.progress((i + 1) / max(1, len(files)), f"Hashed {i + 1}/{len(files)}")
    if p.get("write_manifest") and os.path.isdir(src_abs):
        name = _out_name(src_abs, "SHA256SUMS.txt")
        with open(os.path.join(src_abs, name), "w", encoding="utf-8") as f:
            for k, v in results.items():
                f.write(f"{v}  {k}\n")
    return {"files": len(results), "sha256": results if len(results) <= 2000 else None}


def find_duplicates(root_ids: list[int], rel: str = "", ctx: JobContext | None = None, min_size: int = 1) -> list[dict]:
    by_size: dict[int, list[tuple[int, str, str]]] = defaultdict(list)
    for rid in root_ids:
        root = get_root(rid)
        src_abs, _ = resolve(root, rel, must_exist=True)
        for fp in _walk_files(src_abs):
            if ctx:
                ctx.check()
            try:
                size = os.path.getsize(fp)
            except OSError:
                continue
            if size >= min_size:
                by_size[size].append((rid, fp, os.path.relpath(fp, root.real).replace("\\", "/")))
    groups = []
    candidates = [v for v in by_size.values() if len(v) > 1]
    for i, items in enumerate(candidates):
        by_hash: dict[str, list] = defaultdict(list)
        for rid, fp, r in items:
            if ctx:
                ctx.check()
            try:
                by_hash[fileops.checksum(fp)].append({"root_id": rid, "path": r})
            except OSError:
                continue
        for h, members in by_hash.items():
            if len(members) > 1:
                groups.append({"sha256": h, "size": os.path.getsize(items[0][1]), "files": members})
        if ctx:
            ctx.progress((i + 1) / len(candidates), f"Compared {i + 1}/{len(candidates)} size groups")
    groups.sort(key=lambda g: g["size"] * (len(g["files"]) - 1), reverse=True)
    return groups


@handler("find_duplicates", "Find duplicate files")
def job_dupes(ctx: JobContext, p: dict):
    groups = find_duplicates([int(p["root_id"])], p.get("path", ""), ctx)
    wasted = sum(g["size"] * (len(g["files"]) - 1) for g in groups)
    return {"groups": groups[:500], "group_count": len(groups), "reclaimable_bytes": wasted}


@handler("batch_extract_text", "Extract text from documents")
def job_batch_text(ctx: JobContext, p: dict):
    root = get_root(int(p["root_id"]))
    src_abs, _ = resolve(root, p.get("path", ""), must_exist=True)
    dest_abs, dest_rel = resolve(root, p["dest_dir"], must_exist=True)
    docs = [f for f in _walk_files(src_abs) if os.path.splitext(f)[1].lower() in extract.DOC_EXTS]
    written = []
    for i, fp in enumerate(docs):
        ctx.check()
        try:
            text = extract.extract_text(fp, 5_000_000)
        except Exception as e:
            ctx.log(f"skip {fp}: {e}")
            continue
        name = _out_name(dest_abs, os.path.basename(fp) + ".txt")
        with open(os.path.join(dest_abs, name), "w", encoding="utf-8") as f:
            f.write(text)
        written.append(f"{dest_rel}/{name}" if dest_rel else name)
        ctx.progress((i + 1) / len(docs), f"Extracted {i + 1}/{len(docs)}")
    for rel in written:
        index.upsert(root, resolve(root, rel)[0], rel)
    return {"documents": len(docs), "written": written}


CONVERSIONS = {("xlsx", "csv"), ("csv", "xlsx"), ("docx", "txt"), ("pdf", "txt"), ("json", "csv"), ("tsv", "csv")}


@handler("convert", "Convert file format")
def job_convert(ctx: JobContext, p: dict):
    root = get_root(int(p["root_id"]))
    src_abs, src_rel = resolve(root, p["path"], must_exist=True)
    src_ext = os.path.splitext(src_abs)[1].lower().lstrip(".")
    target = p["target"].lower().lstrip(".")
    if (src_ext, target) not in CONVERSIONS:
        raise ValueError(f"Conversion {src_ext} -> {target} is not supported. Supported: "
                         + ", ".join(f"{a}->{b}" for a, b in sorted(CONVERSIONS)))
    parent = os.path.dirname(src_abs)
    name = _out_name(parent, os.path.splitext(os.path.basename(src_abs))[0] + "." + target)
    out = os.path.join(parent, name)
    if target == "txt":
        with open(out, "w", encoding="utf-8") as f:
            f.write(extract.extract_text(src_abs, 50_000_000))
    elif (src_ext, target) == ("xlsx", "csv"):
        from openpyxl import load_workbook
        wb = load_workbook(src_abs, read_only=True, data_only=True)
        with open(out, "w", newline="", encoding="utf-8-sig") as f:
            w = csv.writer(f)
            for row in wb.worksheets[0].iter_rows(values_only=True):
                w.writerow(["" if v is None else v for v in row])
        wb.close()
    elif target == "xlsx":
        from openpyxl import Workbook
        wb = Workbook(write_only=True)
        ws = wb.create_sheet("Sheet1")
        with open(src_abs, newline="", encoding="utf-8-sig", errors="replace") as f:
            for row in csv.reader(f, delimiter="\t" if src_ext == "tsv" else ","):
                ws.append([_num(v) for v in row])
        wb.save(out)
    elif src_ext == "tsv":
        with open(src_abs, newline="", encoding="utf-8-sig", errors="replace") as f, \
                open(out, "w", newline="", encoding="utf-8-sig") as g:
            csv.writer(g).writerows(csv.reader(f, delimiter="\t"))
    elif src_ext == "json":
        import json
        with open(src_abs, encoding="utf-8") as f:
            data = json.load(f)
        if isinstance(data, dict):
            data = next((v for v in data.values() if isinstance(v, list)), [data])
        keys: list[str] = []
        for item in data:
            if isinstance(item, dict):
                keys += [k for k in item if k not in keys]
        with open(out, "w", newline="", encoding="utf-8-sig") as g:
            w = csv.DictWriter(g, fieldnames=keys or ["value"])
            w.writeheader()
            for item in data:
                w.writerow(item if isinstance(item, dict) else {"value": item})
    rel_parent = os.path.dirname(src_rel)
    rel = f"{rel_parent}/{name}" if rel_parent else name
    index.upsert(root, out, rel)
    return {"output": rel}


def _num(v: str):
    try:
        return int(v)
    except ValueError:
        try:
            return float(v)
        except ValueError:
            return v


@handler("hash_file", "Compute file checksum")
def job_hash(ctx: JobContext, p: dict):
    root = get_root(int(p["root_id"]))
    src_abs, rel = resolve(root, p["path"], must_exist=True)
    h = hashlib.sha256()
    size = os.path.getsize(src_abs) or 1
    done = 0
    with open(src_abs, "rb") as f:
        for chunk in iter(lambda: f.read(4 * 1024 * 1024), b""):
            ctx.check()
            h.update(chunk)
            done += len(chunk)
            ctx.progress(done / size)
    return {"path": rel, "sha256": h.hexdigest(), "computed_at": time.time()}
